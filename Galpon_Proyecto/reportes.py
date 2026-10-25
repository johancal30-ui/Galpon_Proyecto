"""Construye el reporte anual en Excel (HU 11). No depende de Django."""
from io import BytesIO

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Protection, Side
from openpyxl.utils import get_column_letter

MESES = ["Enero", "Febrero", "Marzo", "Abril", "Mayo", "Junio", "Julio",
         "Agosto", "Septiembre", "Octubre", "Noviembre", "Diciembre"]

# Colores
C_BANNER = "F29E4C"     # naranja del home
C_AVISO = "FFFF00"      # amarillo del aviso
C_ENCABEZADO = "D9D9D9" # gris
C_BLOQUEADO = "F2F2F2"  # gris claro: datos y cálculos protegidos
C_TOTAL = "D9D9D9"
C_PROMEDIO = "B6E08A"   # verde
BORDE = Border(*(Side(style="thin", color="808080"),) * 4)
FMT_DINERO = '"$" #,##0'
FMT_ENTERO = '#,##0'


def _encabezado(celda):
    celda.font = Font(bold=True, color="000000")
    celda.fill = PatternFill("solid", fgColor=C_ENCABEZADO)
    celda.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    celda.border = BORDE


def construir_reporte(anio, ventas_mes, gastos_mes, huevos_mes,
                      producciones=(), ventas=(), gastos=(),
                      compra_bultos_mes=None, consumo_bultos_mes=None,
                      porcentaje_ganancia=1 / 3):
    """
    *_mes: dict {mes (1-12): valor}. Si un mes no está en el dict, la fila queda VACÍA
    (así los promedios solo cuentan los meses con datos y los totales no se rompen).
    porcentaje_ganancia: parte de la utilidad que es ganancia del propietario
    (editable dentro del Excel, celda D3).
    Devuelve los bytes del .xlsx.
    """
    compra_bultos_mes = compra_bultos_mes or {}
    consumo_bultos_mes = consumo_bultos_mes or {}

    wb = Workbook()
    ws = wb.active
    ws.title = "Reporte anual"

    # ---- Banner y aviso ----
    ws["A1"] = f"GALPÓN CASTILLO - REPORTE ANUAL {anio}"
    ws["A1"].font = Font(bold=True, size=18, color="3B2412")
    ws["A1"].alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.merge_cells("A1:I1")
    for col in range(1, 10):
        ws.cell(row=1, column=col).fill = PatternFill("solid", fgColor=C_BANNER)
    ws.row_dimensions[1].height = 38

    ws["A2"] = "AVISO: Modificar SOLO las celdas de color BLANCO."
    ws["A2"].font = Font(bold=True, size=11)
    ws["A2"].alignment = Alignment(horizontal="center")
    ws.merge_cells("A2:I2")
    for col in range(1, 10):
        ws.cell(row=2, column=col).fill = PatternFill("solid", fgColor=C_AVISO)

    # ---- % de ganancia (editable) ----
    ws["A3"] = "Porcentaje de la utilidad que es ganancia del propietario:"
    ws["A3"].font = Font(italic=True, size=10)
    ws.merge_cells("A3:C3")
    ws["D3"] = porcentaje_ganancia
    ws["D3"].number_format = "0.00%"
    ws["D3"].font = Font(bold=True)
    ws["D3"].alignment = Alignment(horizontal="center")
    ws["D3"].border = BORDE
    ws["D3"].protection = Protection(locked=False)   # blanca y editable

    # ---- Encabezados (fila 5) ----
    encabezados = ["MES", "VENTAS", "GASTOS", "UTILIDADES NETAS", "GANANCIAS",
                   "COMPRA ALIMENTO / BULTOS", "CONSUMO / BULTOS",
                   "PRODUCCIÓN HUEVOS", "NOTAS"]
    for col, texto in enumerate(encabezados, start=1):
        _encabezado(ws.cell(row=5, column=col, value=texto))
    ws.row_dimensions[5].height = 36

    PRIMERA, ULTIMA = 6, 17          # filas de enero a diciembre
    for i, nombre in enumerate(MESES):
        fila = PRIMERA + i
        mes = i + 1

        def valor(d, convertir):
            v = d.get(mes)
            return convertir(v) if v is not None else None

        ws.cell(row=fila, column=1, value=nombre.upper())
        ws.cell(row=fila, column=2, value=valor(ventas_mes, float))
        ws.cell(row=fila, column=3, value=valor(gastos_mes, float))
        # Utilidades netas = ventas - gastos; vacía si el mes no tiene datos
        ws.cell(row=fila, column=4,
                value=f'=IF(AND(B{fila}="",C{fila}=""),"",N(B{fila})-N(C{fila}))')
        # Ganancias = parte del propietario; 0 si el mes dio pérdida
        ws.cell(row=fila, column=5,
                value=f'=IF(D{fila}="","",MAX(D{fila},0)*$D$3)')
        ws.cell(row=fila, column=6, value=valor(compra_bultos_mes, float))
        ws.cell(row=fila, column=7, value=valor(consumo_bultos_mes, float))
        ws.cell(row=fila, column=8, value=valor(huevos_mes, int))
        ws.cell(row=fila, column=9, value=None)

        for col in range(1, 10):
            c = ws.cell(row=fila, column=col)
            c.border = BORDE
            if col <= 8:   # datos y cálculos: protegidos y en gris claro
                c.fill = PatternFill("solid", fgColor=C_BLOQUEADO)
                c.protection = Protection(locked=True)
            else:          # notas: blanca y editable
                c.protection = Protection(locked=False)
        ws.cell(row=fila, column=1).alignment = Alignment(horizontal="right")
        for col in (2, 3, 4, 5):
            ws.cell(row=fila, column=col).number_format = FMT_DINERO
        for col in (6, 7, 8):
            ws.cell(row=fila, column=col).number_format = FMT_ENTERO
            ws.cell(row=fila, column=col).alignment = Alignment(horizontal="center")

    # ---- Total anual (18) y Promedio mensual (19) ----
    TOTAL, PROMEDIO = ULTIMA + 1, ULTIMA + 2
    ws.cell(row=TOTAL, column=1, value="TOTAL ANUAL")
    ws.cell(row=PROMEDIO, column=1, value="PROMEDIO MENSUAL")
    for col in range(2, 9):
        letra = get_column_letter(col)
        rango = f"{letra}{PRIMERA}:{letra}{ULTIMA}"
        ws.cell(row=TOTAL, column=col, value=f"=SUM({rango})")
        # IFERROR: si ningún mes tiene datos no se rompe, queda en 0
        ws.cell(row=PROMEDIO, column=col, value=f"=IFERROR(AVERAGE({rango}),0)")
    for fila, color in ((TOTAL, C_TOTAL), (PROMEDIO, C_PROMEDIO)):
        for col in range(1, 10):
            c = ws.cell(row=fila, column=col)
            c.font = Font(bold=True)
            c.fill = PatternFill("solid", fgColor=color)
            c.border = BORDE
            c.protection = Protection(locked=True)
        ws.cell(row=fila, column=1).alignment = Alignment(horizontal="right")
        for col in (2, 3, 4, 5):
            ws.cell(row=fila, column=col).number_format = FMT_DINERO
        for col in (6, 7, 8):
            ws.cell(row=fila, column=col).number_format = FMT_ENTERO
            ws.cell(row=fila, column=col).alignment = Alignment(horizontal="center")

    for col, ancho in zip("ABCDEFGHI", (24, 17, 17, 18, 17, 18, 14, 16, 36)):
        ws.column_dimensions[col].width = ancho
    ws.freeze_panes = "A6"

    # ---- Gráfico de barras: ventas, gastos y utilidades por mes ----
    grafico = BarChart()
    grafico.type = "col"
    grafico.title = f"Ventas y utilidades por cada mes del año {anio}"
    grafico.y_axis.title = "Montos"
    grafico.x_axis.title = "Meses"
    grafico.y_axis.number_format = '"$" #,##0'
    for col, color in ((2, "F7C948"), (3, "C0392B"), (4, "6B8E23")):
        grafico.add_data(Reference(ws, min_col=col, min_row=5, max_row=ULTIMA),
                         titles_from_data=True)
        grafico.series[-1].graphicalProperties.solidFill = color
    grafico.set_categories(Reference(ws, min_col=1, min_row=PRIMERA, max_row=ULTIMA))
    grafico.x_axis.tickLblPos = "low"   # nombres de meses abajo, aunque haya pérdidas
    grafico.x_axis.delete = False
    grafico.y_axis.delete = False
    grafico.height = 10
    grafico.width = 26
    ws.add_chart(grafico, f"A{PROMEDIO + 3}")

    # ---- Protección de la hoja (sin contraseña) ----
    ws.protection.sheet = True

    # ---- Hojas de detalle (respaldo de los registros) ----
    def hoja_detalle(titulo, columnas, filas, formatos=()):
        h = wb.create_sheet(titulo)
        for col, texto in enumerate(columnas, start=1):
            _encabezado(h.cell(row=1, column=col, value=texto))
        for r, fila in enumerate(filas, start=2):
            for col, valor in enumerate(fila, start=1):
                c = h.cell(row=r, column=col, value=valor)
                c.border = BORDE
                if col in formatos:
                    c.number_format = FMT_DINERO
        for col in range(1, len(columnas) + 1):
            h.column_dimensions[get_column_letter(col)].width = 20
        h.freeze_panes = "A2"

    hoja_detalle("Producción", ["Fecha", "Huevos", "Rotos", "Cubetas"], producciones)
    hoja_detalle("Ventas", ["Fecha", "Cliente", "Cubetas", "Precio unit.", "Total"],
                 ventas, formatos=(4, 5))
    hoja_detalle("Gastos", ["Fecha", "Categoría", "Descripción", "Monto"],
                 gastos, formatos=(4,))

    salida = BytesIO()
    wb.save(salida)
    return salida.getvalue()    